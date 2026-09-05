"""
Generate the "Top 30 Generative AI Use Cases for Law Firms" report as a Word document.
Requires: pip install python-docx
Run from project root: python scripts/generate_legal_ai_report.py
Output: docs/Top_30_Generative_AI_Use_Cases_Law_Firms.docx
"""
import sys
from pathlib import Path

script_dir = Path(__file__).resolve().parent
if str(script_dir) not in sys.path:
    sys.path.insert(0, str(script_dir))

try:
    from docx import Document
except ImportError:
    print("Install python-docx: pip install python-docx")
    raise

from legal_ai_use_cases_data import USE_CASES


def build_document():
    doc = Document()
    # Title
    doc.add_heading("Top 30 Generative AI Use Cases for Law Firms", level=0)
    doc.add_paragraph(
        "A report on adoption, impact, complexity, value justification, "
        "user stories, user journeys, and personas for generative AI in legal practice."
    )
    doc.add_paragraph()

    # Executive summary
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "This report identifies and ranks 30 high-value use cases for generative AI (GenAI) "
        "in law firms and legal departments. Each use case is evaluated on three dimensions: "
        "Adoption (how widely it is already in use), Impact (business and client value), and "
        "Complexity (implementation difficulty). The ranking from 1 to 30 reflects a combined "
        "view of these factors, with the highest-ranked use cases offering strong adoption and "
        "impact with manageable complexity. For each use case, the report provides exhaustive "
        "detail: value justification for the law firm, user stories, user journey steps, and "
        "personas who would use or benefit from the capability."
    )
    doc.add_paragraph(
        "Sources include Gartner (e.g., top GenAI use cases for legal departments), Thomson Reuters "
        "and Deloitte surveys on legal AI adoption, Clio and legal tech vendor reports, and industry "
        "analysis on ROI and implementation. Adoption rates have more than doubled in recent years, "
        "with contract review, legal research, and document summarization among the most widely "
        "adopted and highest-impact applications."
    )
    doc.add_paragraph()

    # Summary table: Rank | Use Case | Adoption | Impact | Complexity
    doc.add_heading("Summary: Ranking (1–30)", level=1)
    table = doc.add_table(rows=1, cols=5, style="Table Grid")
    h = table.rows[0].cells
    h[0].text = "Rank"
    h[1].text = "Use Case"
    h[2].text = "Adoption (1–10)"
    h[3].text = "Impact (1–10)"
    h[4].text = "Complexity (1–10)"
    for uc in USE_CASES:
        row = table.add_row().cells
        row[0].text = str(uc["rank"])
        row[1].text = uc["name"]
        row[2].text = str(uc["adoption"])
        row[3].text = str(uc["impact"])
        row[4].text = str(uc["complexity"])
    doc.add_paragraph()

    # Detailed section per use case
    doc.add_heading("Detailed Use Case Descriptions", level=1)
    for uc in USE_CASES:
        doc.add_heading(f"{uc['rank']}. {uc['name']}", level=2)

        # Ranking table for this use case
        t = doc.add_table(rows=4, cols=2, style="Table Grid")
        t.rows[0].cells[0].text = "Dimension"
        t.rows[0].cells[1].text = "Score and notes"
        t.rows[1].cells[0].text = "Adoption"
        t.rows[1].cells[1].text = f"{uc['adoption']}/10 — {uc['adoption_note']}"
        t.rows[2].cells[0].text = "Impact"
        t.rows[2].cells[1].text = f"{uc['impact']}/10 — {uc['impact_note']}"
        t.rows[3].cells[0].text = "Complexity"
        t.rows[3].cells[1].text = f"{uc['complexity']}/10 — {uc['complexity_note']}"
        doc.add_paragraph()

        doc.add_heading("Value to the Law Firm", level=3)
        doc.add_paragraph(uc["value"])
        doc.add_paragraph()

        doc.add_heading("User Stories", level=3)
        for story in uc["user_stories"]:
            doc.add_paragraph(story)
        doc.add_paragraph()

        doc.add_heading("User Journey", level=3)
        for i, step in enumerate(uc["journey"], 1):
            doc.add_paragraph(f"{i}. {step}")
        doc.add_paragraph()

        doc.add_heading("Personas", level=3)
        for persona in uc["personas"]:
            doc.add_paragraph(persona)
        doc.add_paragraph()

    # Methodology and sources
    doc.add_heading("Methodology and Sources", level=1)
    doc.add_paragraph(
        "Rankings and narrative are based on published research and industry reports, including: "
        "Gartner (top GenAI use cases for legal departments, 2025); Thomson Reuters (GenAI and "
        "legal professionals, 2025–2026); Deloitte (GenAI use cases for CLOs); Clio and other "
        "legal tech surveys; ABA and law practice reports on AI adoption; and vendor case studies "
        "on contract review, e-discovery, intake, billing, and matter management. Adoption scores "
        "reflect current survey data where available; impact and complexity are assessed from "
        "reported outcomes and implementation descriptions. User stories, journeys, and personas "
        "are synthesized from typical legal workflows and buyer/user research in legal technology."
    )
    doc.add_paragraph()

    doc.add_heading("Document Information", level=1)
    doc.add_paragraph("Report generated for the Legal AI Portal project (003-legal).")
    doc.add_paragraph("Output format: Microsoft Word (.docx).")

    return doc


def main():
    project_root = script_dir.parent
    out_dir = project_root / "docs"
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / "Top_30_Generative_AI_Use_Cases_Law_Firms.docx"
    doc = build_document()
    doc.save(str(out_path))
    print(f"Report saved: {out_path}")


if __name__ == "__main__":
    main()
